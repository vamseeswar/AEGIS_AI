"""Unit Tests for Phase 6: RAG Engine, Multi-Format Parsers, Recursive Chunker, Vector Search & Citation Attributor."""

import io
import json
import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter
from fastapi.testclient import TestClient

from backend.db.base import Base
from backend.db.init_db import initialize_database
from backend.db.session import engine
from backend.main import app
from backend.rag.chunking import RecursiveCharacterChunker
from backend.rag.citation import extract_and_verify_citations, format_context_with_citations
from backend.rag.embeddings import LocalEmbeddingProvider
from backend.rag.parsers import (
    CSVParser,
    DocxParser,
    JSONParser,
    PDFParser,
    ParsedPage,
    TextParser,
)
from backend.rag.retriever import SearchResult, compute_cosine_similarity


@pytest.fixture(scope="module")
def client():
    """Provides a TestClient with clean database state and initialized tables."""
    import asyncio

    async def _reset():
        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.drop_all)
            await conn.run_sync(Base.metadata.create_all)
        await initialize_database()

    asyncio.run(_reset())

    with TestClient(app) as test_client:
        yield test_client


def get_token(client: TestClient, email: str, password: str) -> str:
    res = client.post("/api/v1/auth/login", json={"email": email, "password": password})
    assert res.status_code == 200
    return res.json()["access_token"]


# ------------------------------------------------------------
# 1. Document Parsers Unit Tests
# ------------------------------------------------------------


def test_text_parser():
    parser = TextParser()
    content = b"Line 1: System initialization.\nLine 2: Ready for operations."
    pages = parser.parse(content, "test.txt")
    assert len(pages) == 1
    assert pages[0].page_number == 1
    assert "System initialization" in pages[0].text


def test_csv_parser():
    parser = CSVParser()
    csv_bytes = b"id,metric,value\n1,throughput,450\n2,latency,12"
    pages = parser.parse(csv_bytes, "data.csv")
    assert len(pages) == 1
    assert "Header: id, metric, value" in pages[0].text
    assert "metric=throughput" in pages[0].text


def test_json_parser():
    parser = JSONParser()
    data = {"status": "healthy", "service": "rag_engine", "workers": 4}
    json_bytes = json.dumps(data).encode("utf-8")
    pages = parser.parse(json_bytes, "status.json")
    assert len(pages) == 1
    assert '"status": "healthy"' in pages[0].text


def test_docx_parser():
    parser = DocxParser()
    doc = DocxDocument()
    doc.add_paragraph("Paragraph 1: Autonomous Agent Governance.")
    doc.add_paragraph("Paragraph 2: AST Safety Inspection Protocol.")
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    pages = parser.parse(stream.getvalue(), "test.docx")
    assert len(pages) == 1
    assert "Autonomous Agent Governance" in pages[0].text
    assert "AST Safety Inspection Protocol" in pages[0].text


def test_pdf_parser():
    parser = PDFParser()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    stream = io.BytesIO()
    writer.write(stream)
    stream.seek(0)

    pages = parser.parse(stream.getvalue(), "blank.pdf")
    assert len(pages) >= 1
    assert pages[0].page_number == 1


# ------------------------------------------------------------
# 2. Recursive Chunker Unit Tests
# ------------------------------------------------------------


def test_recursive_chunker_page_preservation():
    chunker = RecursiveCharacterChunker(chunk_size=100, chunk_overlap=20)
    pages = [
        ParsedPage(page_number=1, text="Page one text content that is relatively short."),
        ParsedPage(
            page_number=2,
            text="Page two text content. This page has multiple sentences. We will check that page number 2 is preserved properly.",
        ),
    ]

    chunks = chunker.chunk_pages(pages)
    assert len(chunks) >= 2
    # Ensure page numbers are preserved
    assert any(c.page_number == 1 for c in chunks)
    assert any(c.page_number == 2 for c in chunks)
    assert all(len(c.content) <= 150 for c in chunks)


def test_recursive_chunker_invalid_overlap():
    with pytest.raises(ValueError):
        RecursiveCharacterChunker(chunk_size=100, chunk_overlap=150)


# ------------------------------------------------------------
# 3. Embeddings & Cosine Similarity Tests
# ------------------------------------------------------------


@pytest.mark.asyncio
async def test_local_embedding_provider():
    provider = LocalEmbeddingProvider()
    assert provider.dimension == 384

    texts = ["AEGIS AI autonomous platform", "Cloud database security"]
    vectors = await provider.embed_texts(texts)
    assert len(vectors) == 2
    assert len(vectors[0]) == 384

    # Identical query should have similarity close to 1.0
    query_vec = await provider.embed_query("AEGIS AI autonomous platform")
    sim = compute_cosine_similarity(vectors[0], query_vec)
    assert sim >= 0.99


# ------------------------------------------------------------
# 4. Citation Attribution & Verification Tests
# ------------------------------------------------------------


def test_format_context_and_verify_citations():
    chunks = [
        SearchResult(
            chunk_id="chk-1",
            document_id="doc-12345678",
            document_title="Operations Guide",
            filename="ops.pdf",
            page_number=3,
            content="Standard operating procedure requires human approval for critical writes.",
            score=0.92,
            citation_token="[doc-1234:p.3]",
        ),
        SearchResult(
            chunk_id="chk-2",
            document_id="doc-87654321",
            document_title="Security Policy",
            filename="sec.pdf",
            page_number=7,
            content="All API requests must carry a valid HMAC or JWT authorization header.",
            score=0.88,
            citation_token="[doc-8765:p.7]",
        ),
    ]

    context = format_context_with_citations(chunks)
    assert "<context_documents>" in context
    assert 'token="[doc-1234:p.3]"' in context

    # Test verification
    llm_answer = (
        "Human approval is mandatory for critical operations [doc-1234:p.3], "
        "and auth headers are required [doc-8765:p.7]. Also [fake-999:p.1] was cited."
    )
    verified, unverified = extract_and_verify_citations(llm_answer, chunks)

    assert len(verified) == 2
    assert verified[0].citation_token == "[doc-1234:p.3]"
    assert verified[1].citation_token == "[doc-8765:p.7]"
    assert unverified == ["[fake-999:p.1]"]


# ------------------------------------------------------------
# 5. RAG End-to-End Pipeline & Multi-Tenancy Tests
# ------------------------------------------------------------


def test_rag_ingest_reindex_and_query_flow(client: TestClient):
    # 1. Register User A in Org Alpha
    client.post(
        "/api/v1/auth/register",
        json={
            "email": "rag_alpha@test.com",
            "password": "Passw0rd!",
            "full_name": "Alpha Researcher",
            "organization_name": "OrgAlphaRAG",
        },
    )
    token_a = get_token(client, "rag_alpha@test.com", "Passw0rd!")
    headers_a = {"Authorization": f"Bearer {token_a}"}

    # 2. Upload document as User A
    file_content = (
        b"AEGIS Autonomous AI Platform Architecture.\n\n"
        b"Section 1: Multi-Agent state machine uses LangGraph with supervisor routing.\n\n"
        b"Section 2: AST safety parsing rejects all DML and DDL mutations in SQL analyst agent.\n\n"
        b"Section 3: pgvector stores chunk embeddings with cosine similarity distance."
    )
    up_resp = client.post(
        "/api/v1/documents/upload",
        headers=headers_a,
        files={"file": ("architecture_overview.txt", io.BytesIO(file_content), "text/plain")},
        data={"title": "Architecture Overview"},
    )
    assert up_resp.status_code == 201
    doc_id = up_resp.json()["document"]["id"]

    # 3. Reindex / chunk document
    reindex_resp = client.post(
        f"/api/v1/documents/{doc_id}/reindex",
        headers=headers_a,
    )
    assert reindex_resp.status_code == 200
    chunks_created = reindex_resp.json()["chunks_created"]
    assert chunks_created >= 1

    # 4. Query RAG vector retrieval as User A
    rag_resp = client.post(
        "/api/v1/rag/query",
        headers=headers_a,
        json={"query": "LangGraph supervisor routing", "top_k": 3},
    )
    assert rag_resp.status_code == 200
    rag_data = rag_resp.json()
    assert rag_data["total_retrieved"] >= 1
    assert "context_documents" in rag_data["formatted_context"]

    # 5. Register User B in Org Beta
    client.post(
        "/api/v1/auth/register",
        json={
            "email": "rag_beta@test.com",
            "password": "Passw0rd!",
            "full_name": "Beta Researcher",
            "organization_name": "OrgBetaRAG",
        },
    )
    token_b = get_token(client, "rag_beta@test.com", "Passw0rd!")
    headers_b = {"Authorization": f"Bearer {token_b}"}

    # 6. Verify User B CANNOT retrieve User A's chunks (Tenant Isolation)
    rag_beta_resp = client.post(
        "/api/v1/rag/query",
        headers=headers_b,
        json={"query": "LangGraph supervisor routing", "top_k": 5},
    )
    assert rag_beta_resp.status_code == 200
    # Must be 0 results for Org Beta because Org Beta has no uploaded documents!
    assert rag_beta_resp.json()["total_retrieved"] == 0
